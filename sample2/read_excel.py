from timeit import timeit
from sample_ext_lib import Book, load_workbook as load_workbook_ext
from openpyxl import load_workbook, Workbook

def read_excel_for_python():
    workbook: Workbook = load_workbook("./sample2/sample.xlsx")
    assert workbook.sheetnames == ["シート1"]
    sheet = workbook["シート1"]
    assert sheet["A1"].value == "a1"
    assert sheet.cell(row=1, column=1).value == "a1"
    values = [cell.value for cells in sheet["A1:C2"] for cell in cells]
    assert values == ["a1", "b1", "c1", "a2", "b2", "c2"]
    sheet["A1"] = "new_a1"
    assert sheet["A1"].value == "new_a1"
    sheet["A1"] = "a1"

def read_excel_for_rust():
    book: Book = load_workbook_ext("./sample2/sample.xlsx")
    assert book.sheetnames == ["シート1"]
    sheet = book.get_sheet_by_name("シート1")
    assert sheet["A1"].value == "a1"
    assert sheet.cell(1, 1).value == "a1"
    assert sheet["A1:C2"] == ["a1", "b1", "c1", "a2", "b2", "c2"]
    # sheet["A1"].value = "new_a1"
    sheet.set_value(1, 1, "new_a1")
    assert sheet["A1"].value == "new_a1" # TODO: sheet["A1"].value == "a1"となる
    sheet.set_value(1, 1, "a1")

time_py = timeit(read_excel_for_python, number=100)
time_rust = timeit(read_excel_for_rust, number=100)

print(f"Python: {time_py:.6f} seconds")
print(f"Rust: {time_rust:.6f} seconds")


